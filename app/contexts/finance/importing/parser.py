"""Структурный парсер xlsx → MonthParse (ТЗ §7 шаг 2, ADR-0002 P2).

ЕДИНСТВЕННОЕ место, где живёт openpyxl и знание о форме таблицы (заменяемость ACL,
A6). Парсинг без хардкода номеров строк (позиции плавают по месяцам):
  - колонки — по шапке R1 (доход A–D, расход E–H, дни I–AM);
  - родитель расхода — заливка ячейки E = FFF9CB9C (детект скоупится на колонку E,
    т.к. тем же цветом залита сводная «Дельта» в доходном блоке);
  - подкатегории — последующие незалитые строки до следующего родителя;
  - доходные сводные (Доход/Дельта/Расход) распознаются ПО ИМЕНИ и исключаются;
  - расходный «Итог» — по заливке FFD9D9D9/FFFFE599.
Год — параметр (`year`); месяц — из имени листа.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .errors import ImportIssue, ImportStructureError
from .model import (
    MonthParse,
    ParsedCategory,
    ParsedIncome,
    ParsedParent,
)

# Заливки (ARGB), сверено по реальному файлу 2026-06-07/10.
FILL_PARENT = "FFF9CB9C"  # родитель расхода (в колонке E)
FILL_EXPENSE_TOTAL = {"FFD9D9D9", "FFFFE599"}  # «Итог» расхода / зеркало

# Колонки (1-based) по шапке R1.
COL_INCOME_NAME = 1  # A
COL_INCOME_PLAN = 2  # B
COL_INCOME_FACT = 3  # C
COL_EXPENSE_NAME = 5  # E
COL_EXPENSE_PLAN = 6  # F
COL_EXPENSE_FACT = 7  # G
COL_DAY_FIRST = 9  # I (день 1)
DAYS_IN_MATRIX = 31  # I…AM

# Сводные строки доходного блока — по имени (нормализованному).
INCOME_SUMMARY_NAMES = {"доход", "дельта", "расход"}
EXPENSE_TOTAL_NAMES = {"итог"}

# Месяц по имени листа (рус., родительный/именительный — как в файле).
_MONTHS = {
    "январь": 1, "февраль": 2, "март": 3, "апрель": 4,
    "май": 5, "июнь": 6, "июль": 7, "август": 8,
    "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12,
}

# Классификация источника дохода по нормализованному имени.
_INCOME_KIND = {
    "переход с предыдущего месяца": "carryover",
    "кэшбек": "cashback",
}
_PARTNER_NAMES = {"катя", "сережа", "серёжа"}


def _norm(value) -> str:
    return str(value).strip().lower() if value is not None else ""


def _cell_fill(ws: Worksheet, row: int, col: int) -> str | None:
    cell = ws.cell(row=row, column=col)
    fill = cell.fill
    if fill is None or fill.patternType is None:
        return None
    fg = fill.fgColor
    return fg.rgb if fg is not None else None


def _to_decimal(value, *, row: int, what: str, issues: list[ImportIssue],
                sheet: str) -> Decimal | None:
    """Числовая ячейка → Decimal; пусто → None; нечисло → issue (fail-loud, §13)."""
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        issues.append(ImportIssue(row=row, reason=f"нечисловое значение в «{what}»: {value!r}", sheet=sheet))
        return None


def _month_from_sheet(sheet_name: str, year: int) -> date | None:
    key = _norm(sheet_name)
    month = _MONTHS.get(key)
    if month is None:
        return None
    return date(year, month, 1)


def _income_kind(name: str) -> str:
    n = _norm(name)
    if n in _INCOME_KIND:
        return _INCOME_KIND[n]
    if n in _PARTNER_NAMES:
        return "partner"
    return "other"


def _verify_header(ws: Worksheet, sheet: str) -> None:
    """Шапка R1 должна совпасть с ожидаемой (иначе неузнанная структура, §13)."""
    expected = {
        COL_INCOME_NAME: "доход",
        COL_INCOME_PLAN: "план",
        COL_INCOME_FACT: "факт",
        COL_EXPENSE_NAME: "расходы",
        COL_EXPENSE_PLAN: "план",
        COL_EXPENSE_FACT: "факт",
    }
    for col, exp in expected.items():
        got = _norm(ws.cell(row=1, column=col).value)
        if got != exp:
            raise ImportStructureError(
                f"Лист «{sheet}»: неузнанная шапка в колонке {col} "
                f"(ожидалось «{exp}», получено «{got or '∅'}»)",
                [ImportIssue(row=1, reason="неузнанная шапка", sheet=sheet)],
            )


def _parse_days(ws: Worksheet, row: int, issues: list[ImportIssue],
                sheet: str) -> dict[int, Decimal]:
    days: dict[int, Decimal] = {}
    for d in range(DAYS_IN_MATRIX):
        val = ws.cell(row=row, column=COL_DAY_FIRST + d).value
        dec = _to_decimal(val, row=row, what=f"день {d + 1}", issues=issues, sheet=sheet)
        if dec is not None and dec != 0:
            days[d + 1] = dec
    return days


def _last_data_row(ws: Worksheet) -> int:
    """Последняя строка, где есть имя в A или E (хвост листа — тысячи пустых строк)."""
    last = 1
    for row in range(2, ws.max_row + 1):
        a = ws.cell(row=row, column=COL_INCOME_NAME).value
        e = ws.cell(row=row, column=COL_EXPENSE_NAME).value
        if (a is not None and str(a).strip()) or (e is not None and str(e).strip()):
            last = row
    return last


def _parse_income(ws: Worksheet, last_row: int, issues: list[ImportIssue],
                  sheet: str) -> list[ParsedIncome]:
    """Источники из блока A–D (до первой сводной строки по имени)."""
    out: list[ParsedIncome] = []
    for row in range(2, last_row + 1):
        name = ws.cell(row=row, column=COL_INCOME_NAME).value
        if name is None or not str(name).strip():
            continue
        if _norm(name) in INCOME_SUMMARY_NAMES:
            # Сводная строка (Доход/Дельта/Расход) — не источник (I5). Доходный блок
            # короткий: после первой сводной источников уже нет.
            break
        plan = _to_decimal(ws.cell(row=row, column=COL_INCOME_PLAN).value,
                           row=row, what="доход.план", issues=issues, sheet=sheet)
        fact = _to_decimal(ws.cell(row=row, column=COL_INCOME_FACT).value,
                           row=row, what="доход.факт", issues=issues, sheet=sheet)
        out.append(ParsedIncome(
            name=str(name).strip(), kind=_income_kind(name), plan=plan, fact=fact
        ))
    return out


def _parse_expenses(ws: Worksheet, last_row: int, issues: list[ImportIssue],
                    sheet: str) -> list[ParsedParent]:
    """Дерево расходов из E–H + дни. Родитель — заливка E=FFF9CB9C; «Итог» исключаем."""
    parents: list[ParsedParent] = []
    current: ParsedParent | None = None
    for row in range(2, last_row + 1):
        name = ws.cell(row=row, column=COL_EXPENSE_NAME).value
        if name is None or not str(name).strip():
            continue
        nname = _norm(name)
        fill = _cell_fill(ws, row, COL_EXPENSE_NAME)
        # Расходный «Итог» (по имени или заливке) — производное (I5), пропускаем.
        if nname in EXPENSE_TOTAL_NAMES or fill in FILL_EXPENSE_TOTAL:
            continue
        plan = _to_decimal(ws.cell(row=row, column=COL_EXPENSE_PLAN).value,
                           row=row, what="расход.план", issues=issues, sheet=sheet)
        fact = _to_decimal(ws.cell(row=row, column=COL_EXPENSE_FACT).value,
                           row=row, what="расход.факт", issues=issues, sheet=sheet)
        days = _parse_days(ws, row, issues, sheet)
        if fill == FILL_PARENT:
            current = ParsedParent(
                name=str(name).strip(),
                plan=plan,
                fact=fact if fact is not None else Decimal("0"),
                days=days,
            )
            parents.append(current)
        else:
            if current is None:
                # Подкатегория без родителя — структура не распознана (fail-loud).
                issues.append(ImportIssue(
                    row=row, sheet=sheet,
                    reason=f"подкатегория «{str(name).strip()}» без родителя",
                ))
                continue
            current.children.append(ParsedCategory(
                name=str(name).strip(),
                plan=plan,
                fact=fact if fact is not None else Decimal("0"),
                days=days,
            ))
    return parents


def parse_sheet(ws: Worksheet, year: int) -> MonthParse | None:
    """Один лист → MonthParse. None, если лист не распознан как месяц (пропуск, §13).

    Структурные ошибки (битая шапка, подкатегория без родителя) → ImportStructureError.
    """
    month = _month_from_sheet(ws.title, year)
    if month is None:
        return None  # не месячный лист — пропускаем с пометкой выше
    _verify_header(ws, ws.title)
    issues: list[ImportIssue] = []
    last_row = _last_data_row(ws)
    income = _parse_income(ws, last_row, issues, ws.title)
    parents = _parse_expenses(ws, last_row, issues, ws.title)
    if issues:
        raise ImportStructureError(
            f"Лист «{ws.title}»: не распознано {len(issues)} строк(и)",
            issues,
        )
    if not parents:
        raise ImportStructureError(
            f"Лист «{ws.title}»: не найдено ни одной категории расходов",
            [ImportIssue(row=0, reason="пустой лист расходов", sheet=ws.title)],
        )
    return MonthParse(month=month, sheet_name=ws.title, income=income, parents=parents)


def parse_workbook(data: bytes, year: int) -> tuple[list[MonthParse], list[ImportIssue]]:
    """xlsx-байты → (распознанные месяцы, пропущенные/пустые листы).

    Структурное несоответствие любого месячного листа → ImportStructureError
    (fail-loud, ничего выше не пишется, I3).
    """
    wb = load_workbook(io.BytesIO(data), data_only=True)
    months: list[MonthParse] = []
    skipped: list[ImportIssue] = []
    for ws in wb.worksheets:
        parsed = parse_sheet(ws, year)
        if parsed is None:
            skipped.append(ImportIssue(
                row=0, sheet=ws.title,
                reason=f"лист «{ws.title}» не распознан как месяц — пропущен",
            ))
            continue
        months.append(parsed)
    if not months:
        raise ImportStructureError(
            "В файле не найдено ни одного месячного листа",
            skipped,
        )
    return months, skipped
