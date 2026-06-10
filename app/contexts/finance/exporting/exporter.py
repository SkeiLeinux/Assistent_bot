"""Выгрузка периода в xlsx (ТЗ A7): лист «Сводка» + лист «По дням».

openpyxl заперт здесь (A6) — домен/аналитика его не видят. Данные берутся из
AnalyticsResult (дерево план/факт/разница/%) и сырых дневных строк read-проекции;
выгрузка совпадает с аналитикой (A7). Возвращает bytes готового .xlsx.

«Сводка»: дерево категорий (родитель + дети с отступом), план/факт/разница/%.
«По дням»: матрица подкатегория × дата (только ненулевые дни).
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font

from app.contexts.finance.read import CategoryAggregate
from app.contexts.finance.services.analytics_service import (
    AnalyticsLine,
    AnalyticsResult,
)


def _num(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _flatten(lines: list[AnalyticsLine], depth: int = 0) -> list[tuple[int, AnalyticsLine]]:
    out: list[tuple[int, AnalyticsLine]] = []
    for line in lines:
        out.append((depth, line))
        out.extend(_flatten(line.children, depth + 1))
    return out


def _write_summary(ws, result: AnalyticsResult) -> None:
    ws.title = "Сводка"
    header = ["Категория", "План", "Факт", "Разница", "%"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.append([
        f"Период: {result.date_from.isoformat()} — {result.date_to.isoformat()}",
        None, None, None, None,
    ])
    ws.append([
        "ИТОГО",
        _num(result.total_plan),
        _num(result.total_fact),
        _num(result.total_diff),
        None,
    ])
    for depth, line in _flatten(result.lines):
        ws.append([
            ("    " * depth) + line.name,
            _num(line.plan),
            _num(line.fact),
            _num(line.diff),
            round(line.pct, 4) if line.pct is not None else None,
        ])


def _write_daily(ws, daily: dict[int, dict[date, Decimal]],
                 cat_names: dict[int, str], dates: list[date]) -> None:
    """Матрица подкатегория × дата (только листья с фактом)."""
    header = ["Подкатегория"] + [d.isoformat() for d in dates]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for cat_id in sorted(daily, key=lambda cid: cat_names.get(cid, "")):
        row = [cat_names.get(cat_id, str(cat_id))]
        per_day = daily[cat_id]
        for d in dates:
            amount = per_day.get(d)
            row.append(float(amount) if amount is not None else None)
        ws.append(row)


def export_period(
    result: AnalyticsResult,
    categories: list[CategoryAggregate],
    daily_rows: list[tuple[int, date, Decimal]],
) -> bytes:
    """Собрать книгу выгрузки. `daily_rows` = (category_id, date, amount) листьев."""
    wb = Workbook()
    _write_summary(wb.active, result)

    cat_names = {c.category_id: c.name for c in categories}
    daily: dict[int, dict[date, Decimal]] = {}
    all_dates: set[date] = set()
    for cat_id, day, amount in daily_rows:
        daily.setdefault(cat_id, {})[day] = amount
        all_dates.add(day)
    ws_daily = wb.create_sheet("По дням")
    _write_daily(ws_daily, daily, cat_names, sorted(all_dates))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
