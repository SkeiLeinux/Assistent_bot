"""S2b Finance — аналитика план/факт/иерархия/каденс, доходы, ревью, выгрузка
(критерии A3-resolve/A5/A7 ТЗ spec-finance §15).

Каденс D/W/M/Q/Y + произвольный период; I9 (plan=None вне целых месяцев),
I10 (доходы только целыми месяцами). Сводные (Доход/Расход/Дельта) выводятся (I5).
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

import app.contexts.finance.models  # noqa: F401 — регистрирует finance-таблицы
from app.kernel.models import Family, Membership, User
from app.kernel.sharing import Viewer

from app.contexts.finance.importing import import_xlsx
from app.contexts.finance.models import CategoryStatus
from app.contexts.finance.repositories import CategoryRepository
from app.contexts.finance.services import (
    AnalyticsService,
    InvalidResolution,
    TaxonomyService,
    resolve_period,
)

REAL_FILE = Path("docs/Бюджет общий 2026.xlsx")
TOL = Decimal("1.00")


async def _setup(session):
    u1 = User(telegram_user_id=1, first_name="Катя")
    u2 = User(telegram_user_id=2, first_name="Сережа")
    session.add_all([u1, u2])
    await session.flush()
    fam = Family(created_by=u1.id)
    session.add(fam)
    await session.flush()
    session.add_all([
        Membership(user_id=u1.id, family_id=fam.id),
        Membership(user_id=u2.id, family_id=fam.id),
    ])
    await session.flush()
    viewer = Viewer(id=u1.id, family_id=fam.id)
    await import_xlsx(session, fam.id, REAL_FILE.read_bytes(), 2026)
    await session.commit()
    return fam, viewer


# ── A5: план/факт/разница/% + drill-down по иерархии ──────────────────────────

async def test_analytics_plan_fact_diff_pct_month(session):
    """A5: за целый месяц — план/факт/разница/% по категориям + иерархия."""
    fam, viewer = await _setup(session)
    svc = AnalyticsService(session)
    result = await svc.analytics(viewer, resolve_period("M", month=date(2026, 3, 1)))

    assert result.plan_applicable is True
    # Итог совпадает с «Итог» расхода листа (143586.24) и планом 111000.
    assert abs(result.total_fact - Decimal("143586.24")) <= TOL
    assert abs(result.total_plan - Decimal("111000")) <= TOL
    assert abs(result.total_diff - (result.total_plan - result.total_fact)) <= TOL

    edovoe = next(l for l in result.lines if l.name == "Едовое")
    assert edovoe.children  # drill-down: подкатегории присутствуют
    assert edovoe.diff == edovoe.plan - edovoe.fact
    assert edovoe.pct == pytest.approx(float(edovoe.fact / edovoe.plan))
    # Факт родителя = Σ факт детей (агрегат, I3/I5).
    sum_children = sum((c.fact for c in edovoe.children), Decimal("0"))
    assert abs(edovoe.fact - sum_children) <= TOL


async def test_analytics_drilldown_subtree(session):
    """A5: category_id → поддерево конкретной категории (drill-down)."""
    fam, viewer = await _setup(session)
    svc = AnalyticsService(session)
    full = await svc.analytics(viewer, resolve_period("M", month=date(2026, 3, 1)))
    edovoe = next(l for l in full.lines if l.name == "Едовое")

    sub = await svc.analytics(
        viewer, resolve_period("M", month=date(2026, 3, 1)), category_id=edovoe.category_id
    )
    assert len(sub.lines) == 1
    assert sub.lines[0].category_id == edovoe.category_id
    assert {c.name for c in sub.lines[0].children} == {c.name for c in edovoe.children}


# ── A5 (каденс): D/W/M/Q/Y + произвольный; I9 plan=None вне целых месяцев ──────

async def test_quarter_aggregates_months(session):
    """A5: квартал = агрегация месячных листов; план применим (целые месяцы)."""
    fam, viewer = await _setup(session)
    svc = AnalyticsService(session)
    q = await svc.analytics(viewer, resolve_period("Q", year=2026, quarter=2))
    # Q2 = апрель+май+июнь; план применим (целые месяцы, I9).
    assert q.plan_applicable is True
    assert q.total_plan is not None
    # Факт Q2 > 0 (апрель/май с фактом; июнь нулевой).
    assert q.total_fact > 0


async def test_partial_period_has_no_plan(session):
    """A5/I9: D/W/частичный диапазон → plan=None (UI показывает только факт)."""
    fam, viewer = await _setup(session)
    svc = AnalyticsService(session)
    d = await svc.analytics(
        viewer, resolve_period("D", date_from=date(2026, 3, 1), date_to=date(2026, 3, 10))
    )
    assert d.plan_applicable is False
    assert d.total_plan is None
    assert d.total_diff is None
    assert all(l.plan is None and l.diff is None for l in d.lines)
    # Факт за частичный период доступен и > 0.
    assert d.total_fact > 0


async def test_arbitrary_range_by_days(session):
    """A5: произвольный период считается по дневным строкам (S5)."""
    fam, viewer = await _setup(session)
    svc = AnalyticsService(session)
    full = await svc.analytics(
        viewer, resolve_period("custom", date_from=date(2026, 3, 1), date_to=date(2026, 3, 31))
    )
    half = await svc.analytics(
        viewer, resolve_period("custom", date_from=date(2026, 3, 1), date_to=date(2026, 3, 15))
    )
    # Часть месяца ≤ весь месяц по факту.
    assert half.total_fact <= full.total_fact
    assert half.plan_applicable is False  # не целый месяц (I9)


# ── A5 (доходы): партнёры + переходящий остаток + Дельта; I10 ─────────────────

async def test_income_report_whole_month(session):
    """A5: доходы по источникам/партнёрам + carryover + Дельта = доход−расход (I5)."""
    fam, viewer = await _setup(session)
    svc = AnalyticsService(session)
    report = await svc.income_report(viewer, resolve_period("M", month=date(2026, 3, 1)))

    # Совпадает со сводными строками листа: Доход 144328.31, Дельта 742.07.
    assert abs(report.total_income - Decimal("144328.31")) <= TOL
    assert abs(report.delta - Decimal("742.07")) <= TOL
    assert abs(report.delta - (report.total_income - report.total_expense)) <= TOL
    # Переходящий остаток (carryover) присутствует.
    assert report.carryover == Decimal("92.31")
    kinds = {s.kind for s in report.by_source}
    assert "partner" in kinds and "carryover" in kinds


async def test_income_excluded_for_partial_period(session):
    """A5/I10: период без целого месяца → доходная секция пуста."""
    fam, viewer = await _setup(session)
    svc = AnalyticsService(session)
    report = await svc.income_report(
        viewer, resolve_period("custom", date_from=date(2026, 3, 5), date_to=date(2026, 3, 20))
    )
    assert report.by_source == []
    assert report.total_income == Decimal("0")


# ── A3: ревью → accept / alias переводит в ACTIVE ─────────────────────────────

async def test_resolve_accept_activates(session):
    """A3: accept переводит NEEDS_REVIEW → ACTIVE."""
    fam, viewer = await _setup(session)
    tax = TaxonomyService(session)
    review = await tax.review_queue(viewer)
    target = review[0]
    resolved = await tax.accept(viewer, target.id)
    await session.commit()
    assert resolved.status == CategoryStatus.ACTIVE


async def test_resolve_alias_merges_and_records(session):
    """A3: alias связывает дубль с каноном (CategoryAlias), переносит данные, не fuzzy."""
    fam, viewer = await _setup(session)
    repo = CategoryRepository(session)
    tax = TaxonomyService(session)

    # Возьмём два корня; алиасим один в другой (один уровень — I1).
    roots = [c for c in await repo.list_all(fam.id) if c.parent_id is None]
    canon, dup = roots[0], roots[1]
    # Очистим детей дубля для чистоты (alias запрещает дубль с детьми).
    for child in await repo.children_of(dup.id):
        await session.delete(child)
    await session.flush()

    result = await tax.alias(viewer, dup.id, canon.id)
    await session.commit()
    assert result.id == canon.id
    # Дубль удалён, алиас зарегистрирован.
    assert await repo.get(fam.id, dup.id) is None
    alias = await repo.find_alias(fam.id, dup.name.lower())
    assert alias is not None and alias.category_id == canon.id


async def test_alias_rejects_cross_level(session):
    """A3: алиас между разными уровнями дерева запрещён (I1)."""
    fam, viewer = await _setup(session)
    repo = CategoryRepository(session)
    tax = TaxonomyService(session)
    all_cats = await repo.list_all(fam.id)
    child = next(c for c in all_cats if c.parent_id is not None)
    root = await repo.get(fam.id, child.parent_id)
    with pytest.raises(InvalidResolution):
        await tax.alias(viewer, child.id, root.id)


# ── A7: выгрузка периода в xlsx (Q1-Б1) ───────────────────────────────────────

async def test_export_xlsx_two_sheets_matches_analytics(session):
    """A7: GET /finance/export отдаёт xlsx (Сводка+По дням), данные = аналитике."""
    from openpyxl import load_workbook

    fam, viewer = await _setup(session)
    svc = AnalyticsService(session)
    period = resolve_period("M", month=date(2026, 3, 1))
    data = await svc.export_xlsx(viewer, period)

    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Сводка", "По дням"]

    ws = wb["Сводка"]
    assert [c.value for c in ws[1]] == ["Категория", "План", "Факт", "Разница", "%"]
    # Строка ИТОГО (3-я) совпадает с аналитикой.
    result = await svc.analytics(viewer, period)
    itogo = [ws.cell(row=3, column=c).value for c in range(1, 5)]
    assert itogo[0] == "ИТОГО"
    assert abs(Decimal(str(itogo[2])) - result.total_fact) <= TOL
    assert abs(Decimal(str(itogo[1])) - result.total_plan) <= TOL

    # Лист «По дням» — матрица с датами.
    wd = wb["По дням"]
    assert wd.max_row > 1 and wd.max_column > 1
    assert wd.cell(row=1, column=1).value == "Подкатегория"
