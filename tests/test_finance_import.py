"""S2b Finance — Import ACL: round-trip реального xlsx, идемпотентность, fail-loud,
NEEDS_REVIEW (критерии A1/A2/A3/A4 ТЗ spec-finance §15).

Тесты на сервисном/ACL-слое; домен Telegram-агностичен (принимает Viewer + bytes).
Регистрация finance-таблиц — импортом моделей здесь (conftest их не знает).
A6 (изоляция openpyxl) проверяется import-linter, см. test_finance_isolation.py.
"""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

import app.contexts.finance.models  # noqa: F401 — регистрирует finance-таблицы
from app.kernel.models import Family, Membership, User
from app.kernel.sharing import Viewer

from app.contexts.finance.importing import import_xlsx
from app.contexts.finance.importing.errors import ImportStructureError
from app.contexts.finance.importing.parser import parse_workbook
from app.contexts.finance.importing.validate import validate_all
from app.contexts.finance.models import (
    ActualAmount,
    BudgetPlan,
    Category,
    CategoryStatus,
    Income,
    IncomeSource,
)
from app.contexts.finance.repositories import CategoryRepository

REAL_FILE = Path("docs/Бюджет общий 2026.xlsx")
TOL = Decimal("1.00")


async def _family(session):
    u1 = User(telegram_user_id=1, first_name="Катя")
    u2 = User(telegram_user_id=2, first_name="Сережа")
    session.add_all([u1, u2])
    await session.flush()
    fam = Family(created_by=u1.id)
    session.add(fam)
    await session.flush()
    session.add_all([
        Membership(user_id=u1.id, family_id=fam.id),
        Membership(user_id=u2.id, family_id=fam.id),
    ])
    await session.flush()
    return u1, u2, fam, Viewer(id=u1.id, family_id=fam.id)


def _real_bytes() -> bytes:
    return REAL_FILE.read_bytes()


async def _count(session, model) -> int:
    from sqlalchemy import func, select

    return await session.scalar(select(func.count()).select_from(model))


# ── A1: импорт Март–Июнь, инвариант родитель=Σ детей (факт + дни) ──────────────

async def test_real_file_invariant_holds():
    """A1: парсинг реального файла и валидация инварианта на всех 4 листах.

    Парковка Q-finance-1: инвариант = (факт родителя=Σ факт детей) ∧
    (факт ребёнка=Σ его дней), допуск ≤1 ₽. Дневная строка родителя в источнике
    неполна (Май), но не хранится (родитель-агрегат) — не валидируется."""
    months, skipped = parse_workbook(_real_bytes(), 2026)
    assert [m.sheet_name for m in months] == ["Март", "Апрель", "Май", "Июнь"]
    assert skipped == []
    # Не должно бросить — инвариант держится на всех листах.
    validate_all(months)
    # Явная проверка инварианта (дублирует валидатор, но как контракт A1):
    for m in months:
        for parent in m.parents:
            sum_children = sum((c.fact for c in parent.children), Decimal("0"))
            assert abs(parent.fact - sum_children) <= TOL, (m.sheet_name, parent.name)
            for child in parent.children:
                sum_days = sum(child.days.values(), Decimal("0"))
                assert abs(child.fact - sum_days) <= TOL, (parent.name, child.name)


async def test_real_file_roundtrip_to_domain(session):
    """A1: импорт реального файла → дерево категорий, план, дневной факт, доходы."""
    _, _, fam, viewer = await _family(session)
    outcome, skipped = await import_xlsx(session, fam.id, _real_bytes(), 2026)
    await session.commit()

    assert len(outcome.months) == 4
    assert {m.sheet_name for m in outcome.months} == {"Март", "Апрель", "Май", "Июнь"}

    # Дерево: 10 родителей (одинаковый состав по месяцам), 2 уровня (I1).
    repo = CategoryRepository(session)
    roots = [c for c in await repo.list_all(fam.id) if c.parent_id is None]
    assert len(roots) == 10
    assert {r.name for r in roots} >= {"Едовое", "Тренировки", "Прочее"}
    for c in await repo.list_all(fam.id):
        if c.parent_id is not None:
            parent = await repo.get(fam.id, c.parent_id)
            assert parent.parent_id is None  # глубина ровно 2 (I1)

    # Факт хранится на листьях; родитель = агрегат → у родителя ActualAmount нет.
    from sqlalchemy import select

    root_ids = [r.id for r in roots]
    parent_actuals = await session.execute(
        select(ActualAmount).where(ActualAmount.category_id.in_(root_ids))
    )
    assert parent_actuals.first() is None

    assert await _count(session, ActualAmount) > 0
    assert await _count(session, BudgetPlan) > 0
    assert await _count(session, Income) > 0
    assert await _count(session, IncomeSource) == 5  # 5 источников дохода


async def test_march_parent_fact_equals_sum_children_in_domain(session):
    """A1: после импорта факт родителя (агрегат) == Σ факт детей за март."""
    _, _, fam, viewer = await _family(session)
    await import_xlsx(session, fam.id, _real_bytes(), 2026)
    await session.commit()

    from app.contexts.finance.read import finance_aggregate
    from app.contexts.finance.services.periods import month_period, whole_months

    p = month_period(date(2026, 3, 1))
    res = await finance_aggregate(
        session, fam.id, p.date_from, p.date_to,
        plan_applicable=True, whole_months=whole_months(p.date_from, p.date_to),
    )
    by_id = {a.category_id: a for a in res.categories}
    children_of: dict[int, list] = {}
    for a in res.categories:
        if a.parent_id is not None:
            children_of.setdefault(a.parent_id, []).append(a)
    for a in res.categories:
        if a.parent_id is None and a.category_id in children_of:
            sum_children = sum(
                (c.fact for c in children_of[a.category_id]), Decimal("0")
            )
            assert abs(a.fact - sum_children) <= TOL, a.name


# ── A2: идемпотентный повторный импорт месяца (I6) ─────────────────────────────

async def test_reimport_is_idempotent(session):
    """A2: повторный импорт того же файла не плодит дублей (I6)."""
    _, _, fam, viewer = await _family(session)
    await import_xlsx(session, fam.id, _real_bytes(), 2026)
    await session.commit()
    cats1 = await _count(session, Category)
    acts1 = await _count(session, ActualAmount)
    plans1 = await _count(session, BudgetPlan)
    incs1 = await _count(session, Income)

    outcome2, _ = await import_xlsx(session, fam.id, _real_bytes(), 2026)
    await session.commit()

    assert await _count(session, Category) == cats1
    assert await _count(session, ActualAmount) == acts1
    assert await _count(session, BudgetPlan) == plans1
    assert await _count(session, Income) == incs1
    # Категории уже существуют → ничего нового на ревью.
    assert outcome2.created_for_review == []


# ── A4: структурное несоответствие → fail-loud, ничего не записано (I3) ────────

def _corrupt_break_invariant(data: bytes) -> bytes:
    """Подпорченная копия: ломаем инвариант родитель=Σ детей в марте (G родителя)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    ws = wb["Март"]
    # R2 = «Едовое» (родитель), колонка G(7) = факт. Сильно расходим с Σ детей.
    ws.cell(row=2, column=7).value = 999999
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _corrupt_break_header(data: bytes) -> bytes:
    """Подпорченная копия: ломаем шапку (неузнанная структура)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    wb["Март"].cell(row=1, column=5).value = "ШАПКА-МУСОР"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def test_broken_invariant_fails_loud_without_write(session):
    """A4: расхождение родитель≠Σ → ImportStructureError, ничего не записано (I3)."""
    _, _, fam, viewer = await _family(session)
    broken = _corrupt_break_invariant(_real_bytes())
    with pytest.raises(ImportStructureError) as exc:
        await import_xlsx(session, fam.id, broken, 2026)
    await session.rollback()
    assert exc.value.issues  # отчёт о нераспознанном
    # Ничего не записано (fail-loud до маппера / откат).
    assert await _count(session, Category) == 0
    assert await _count(session, ActualAmount) == 0
    assert await _count(session, Income) == 0


async def test_broken_header_fails_loud(session):
    """A4: неузнанная шапка → ImportStructureError, ничего не записано."""
    _, _, fam, viewer = await _family(session)
    broken = _corrupt_break_header(_real_bytes())
    with pytest.raises(ImportStructureError):
        await import_xlsx(session, fam.id, broken, 2026)
    await session.rollback()
    assert await _count(session, Category) == 0


# ── A3: невиданная категория → NEEDS_REVIEW; alias/accept → ACTIVE ─────────────

async def test_unseen_categories_marked_needs_review(session):
    """A3: первый импорт — все категории невиданные → NEEDS_REVIEW (I7)."""
    _, _, fam, viewer = await _family(session)
    outcome, _ = await import_xlsx(session, fam.id, _real_bytes(), 2026)
    await session.commit()

    review = await CategoryRepository(session).list_review(fam.id)
    total = await CategoryRepository(session).list_all(fam.id)
    # Никаких алиасов/существующих — всё на ревью.
    assert len(review) == len(total)
    assert all(c.status == CategoryStatus.NEEDS_REVIEW for c in review)
    assert outcome.created_for_review  # отражены в отчёте импорта


async def test_year_parameter_applied(session):
    """Год импорта — параметр (Q1-А1): месяцы привязаны к переданному году."""
    _, _, fam, viewer = await _family(session)
    await import_xlsx(session, fam.id, _real_bytes(), 2030)
    await session.commit()

    from sqlalchemy import select

    months = (
        await session.execute(select(ActualAmount.date))
    ).scalars().all()
    assert months and all(d.year == 2030 for d in months)
